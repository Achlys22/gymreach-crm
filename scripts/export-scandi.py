#!/usr/bin/env python3
"""Export Scandinavian gym leads to xlsx + txt"""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

leads = json.load(open('/tmp/scandi_gym_leads.json'))
leads.sort(key=lambda x: (x.get('country', ''), x.get('name', '')))

wb = Workbook()
wb.properties.creator = 'GymReach Hunter'

# Main sheet
ws = wb.active
ws.title = 'Scandinavian Gym Leads'

headers = ['Country', 'Gym Name', 'Instagram Handle', 'Instagram URL', 'City', 'Discipline', 'Source', 'Status']
for i, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=i)
    cell.value = h
    cell.font = Font(bold=True, size=12, color='FFFFFF')
    cell.fill = PatternFill(start_color='F43F5E', end_color='F43F5E', fill_type='solid')

country_colors = {
    'Denmark': PatternFill(start_color='E8F5E9', end_color='E8F5E9', fill_type='solid'),
    'Norway': PatternFill(start_color='E3F2FD', end_color='E3F2FD', fill_type='solid'),
    'Sweden': PatternFill(start_color='FFF3E0', end_color='FFF3E0', fill_type='solid'),
}

for idx, lead in enumerate(leads, 2):
    ws.cell(row=idx, column=1, value=lead.get('country', ''))
    ws.cell(row=idx, column=2, value=lead.get('name', ''))
    ws.cell(row=idx, column=3, value='@' + lead.get('handle', ''))
    ws.cell(row=idx, column=4, value=f"https://instagram.com/{lead.get('handle','')}")
    ws.cell(row=idx, column=5, value=lead.get('city') or '')
    ws.cell(row=idx, column=6, value=lead.get('discipline') or '')
    ws.cell(row=idx, column=7, value=lead.get('source', ''))
    ws.cell(row=idx, column=8, value='Verified' if 'instagram.com' in lead.get('source', '') or lead.get('source') == 'web' else 'To verify')
    
    fill = country_colors.get(lead.get('country', ''))
    if fill:
        ws.cell(row=idx, column=1).fill = fill

# Column widths
widths = {'A': 12, 'B': 30, 'C': 25, 'D': 45, 'E': 18, 'F': 18, 'G': 10, 'H': 12}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

ws.freeze_panes = 'A2'

# Summary sheet
summary = wb.create_sheet('Summary')
summary.cell(row=1, column=1, value='Scandinavian Gym Leads Summary')
summary.cell(row=1, column=1).font = Font(bold=True, size=16)

summary.cell(row=3, column=1, value='Country').font = Font(bold=True)
summary.cell(row=3, column=2, value='Total Leads').font = Font(bold=True)
summary.cell(row=3, column=3, value='Verified (web)').font = Font(bold=True)
summary.cell(row=3, column=4, value='To Verify (LLM)').font = Font(bold=True)

by_country = {}
for l in leads:
    c = l.get('country', 'Unknown')
    if c not in by_country:
        by_country[c] = {'total': 0, 'web': 0, 'llm': 0}
    by_country[c]['total'] += 1
    if 'instagram.com' in l.get('source', '') or l.get('source') == 'web':
        by_country[c]['web'] += 1
    else:
        by_country[c]['llm'] += 1

row = 4
for country in ['Denmark', 'Norway', 'Sweden']:
    c = by_country.get(country, {'total': 0, 'web': 0, 'llm': 0})
    summary.cell(row=row, column=1, value=country)
    summary.cell(row=row, column=2, value=c['total'])
    summary.cell(row=row, column=3, value=c['web'])
    summary.cell(row=row, column=4, value=c['llm'])
    row += 1

summary.cell(row=row, column=1, value='TOTAL').font = Font(bold=True)
summary.cell(row=row, column=2, value=len(leads)).font = Font(bold=True)

summary.column_dimensions['A'].width = 15
summary.column_dimensions['B'].width = 15
summary.column_dimensions['C'].width = 20
summary.column_dimensions['D'].width = 20

wb.save('/home/z/my-project/download/scandinavian-gym-leads.xlsx')
print(f'XLSX saved with {len(leads)} leads')

# Create txt file
txt_lines = []
txt_lines.append('SCANDINAVIAN GYM LEADS')
txt_lines.append('MMA / Muay Thai / Boxing / BJJ / Kickboxing')
txt_lines.append(f'Total: {len(leads)} leads')
txt_lines.append(f'Generated: 2026-08-02')
txt_lines.append('=' * 60)
txt_lines.append('')

current = ''
for lead in leads:
    c = lead.get('country', '')
    if c != current:
        current = c
        count = by_country.get(c, {}).get('total', 0)
        txt_lines.append('')
        txt_lines.append('=' * 60)
        txt_lines.append(f'{current.upper()} ({count} leads)')
        txt_lines.append('=' * 60)
        txt_lines.append('')
    txt_lines.append(lead.get('name', ''))
    txt_lines.append(f"  IG: @{lead.get('handle', '')}")
    txt_lines.append(f"  URL: https://instagram.com/{lead.get('handle', '')}")
    if lead.get('city'):
        txt_lines.append(f"  City: {lead.get('city')}")
    if lead.get('discipline'):
        txt_lines.append(f"  Discipline: {lead.get('discipline')}")
    txt_lines.append(f"  Source: {lead.get('source', '')}")
    txt_lines.append('')

with open('/home/z/my-project/download/scandinavian-gym-leads.txt', 'w') as f:
    f.write('\n'.join(txt_lines))

print('TXT saved')
print('')
print('Summary:')
for country in ['Denmark', 'Norway', 'Sweden']:
    c = by_country.get(country, {'total': 0, 'web': 0, 'llm': 0})
    print(f'  {country}: {c["total"]} ({c["web"]} verified, {c["llm"]} to verify)')
print(f'  TOTAL: {len(leads)}')
